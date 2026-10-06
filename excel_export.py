import os
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from datetime import datetime
import aiosqlite
from config import DB_NAME

async def generate_stock_excel(output_filename: str = "ombor_hisoboti.xlsx") -> str:
    """
    Omborxona qoldiqlari, bo'limlar sarfi va tranzaksiyalar jurnalidan professional Excel fayl yaratadi.
    """
    wb = openpyxl.Workbook()

    # Stillar
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    dept_header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    warning_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    warning_font = Font(name="Calibri", size=11, color="9C0006", bold=True)
    normal_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    normal_font = Font(name="Calibri", size=11, color="006100")
    
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # ---------------- 1-VARAQ: DORILAR QOLDIG'I ----------------
    ws1 = wb.active
    ws1.title = "Dorilar qoldig'i"

    ws1.merge_cells("A1:H1")
    t1 = ws1["A1"]
    t1.value = f"KLINIKA OMBORXONASI DORILAR QOLDIG'I ({datetime.now().strftime('%d.%m.%Y %H:%M')})"
    t1.font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    t1.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 30

    headers1 = [
        "№", "Mahsulot nomi", "Toifasi", "O'lchov birligi", "Mavjud qoldiq", 
        "Min. me'yor", "Joylashuvi", "Yaroqlilik muddati", "Holati"
    ]
    ws1.append([])
    ws1.append(headers1)
    ws1.row_dimensions[3].height = 24

    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    async with aiosqlite.connect(DB_NAME) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute("SELECT * FROM medicines ORDER BY category ASC, name ASC") as cursor:
            medicines = await cursor.fetchall()

    row_num = 4
    cat_names = {"dori": "Dori-darmon", "operatsion": "Operatsion rasxod", "xojalik": "Xo'jalik xarajati"}
    for idx, med in enumerate(medicines, 1):
        is_low = med["quantity"] <= med["min_quantity"]
        status_text = "⚠️ Kam qolgan" if is_low else "Yetarli"

        ws1.append([
            idx,
            med["name"],
            cat_names.get(med["category"], med["category"]),
            med["unit"],
            med["quantity"],
            med["min_quantity"],
            med["location"] or "-",
            med["expiry_date"] or "-",
            status_text
        ])

        status_cell = ws1.cell(row=row_num, column=9)
        qty_cell = ws1.cell(row=row_num, column=5)
        
        if is_low:
            status_cell.fill = warning_fill
            status_cell.font = warning_font
            qty_cell.font = Font(name="Calibri", size=11, bold=True, color="C00000")
        else:
            status_cell.fill = normal_fill
            status_cell.font = normal_font

        for c in range(1, 10):
            ws1.cell(row=row_num, column=c).border = thin_border
            if c in (1, 3, 4, 5, 6, 8, 9):
                ws1.cell(row=row_num, column=c).alignment = Alignment(horizontal="center")

        row_num += 1

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # ---------------- 2-VARAQ: BO'LIMLAR SARFI ----------------
    ws2 = wb.create_sheet(title="Bo'limlar sarfi")
    ws2.merge_cells("A1:E1")
    t2 = ws2["A1"]
    t2.value = f"BO'LIMLAR KESIMIDA CHIQIMLAR HISOBOTI ({datetime.now().strftime('%d.%m.%Y')})"
    t2.font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    t2.alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 30

    headers_dept = ["№", "Bo'lim nomi", "Chiqimlar soni", "Jami olingan miqdor", "Ulush (%)"]
    ws2.append([])
    ws2.append(headers_dept)
    ws2.row_dimensions[3].height = 24

    for col_idx in range(1, len(headers_dept) + 1):
        cell = ws2.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = dept_header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    async with aiosqlite.connect(DB_NAME) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute("""
            SELECT department, COUNT(*) as tx_count, SUM(quantity) as total_qty
            FROM transactions
            WHERE type = 'chiqim' AND department != ''
            GROUP BY department
            ORDER BY total_qty DESC
        """) as cursor:
            dept_stats = await cursor.fetchall()

    total_dept_qty = sum(r["total_qty"] or 0 for r in dept_stats)
    row_num_dept = 4
    for idx, d_stat in enumerate(dept_stats, 1):
        d_qty = d_stat["total_qty"] or 0
        pct = f"{(d_qty / total_dept_qty * 100):.1f}%" if total_dept_qty > 0 else "0%"
        ws2.append([
            idx,
            d_stat["department"],
            d_stat["tx_count"],
            d_qty,
            pct
        ])
        for c in range(1, 6):
            ws2.cell(row=row_num_dept, column=c).border = thin_border
            if c in (1, 3, 4, 5):
                ws2.cell(row=row_num_dept, column=c).alignment = Alignment(horizontal="center")
        row_num_dept += 1

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 4, 15)

    # ---------------- 3-VARAQ: HARAKATLAR JURNALI ----------------
    ws3 = wb.create_sheet(title="Harakatlar jurnali")
    ws3.merge_cells("A1:H1")
    t3 = ws3["A1"]
    t3.value = f"KIRIM VA CHIQIM HARAKATLARI TARIXI ({datetime.now().strftime('%d.%m.%Y')})"
    t3.font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
    t3.alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 30

    headers3 = [
        "Sana va Vaqt", "Amal turi", "Dori nomi", 
        "Miqdori", "Bo'lim", "Izoh / Qabul qiluvchi", "Mas'ul xodim", "ID"
    ]
    ws3.append([])
    ws3.append(headers3)
    ws3.row_dimensions[3].height = 24

    for col_idx in range(1, len(headers3) + 1):
        cell = ws3.cell(row=3, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    async with aiosqlite.connect(DB_NAME) as db:
        db.row_factory = aiosqlite.Row
        async with db.execute("""
            SELECT t.*, m.name as med_name, m.unit as med_unit 
            FROM transactions t 
            JOIN medicines m ON t.medicine_id = m.id 
            ORDER BY t.created_at DESC 
            LIMIT 1000
        """) as cursor:
            transactions = await cursor.fetchall()

    row_num3 = 4
    for tr in transactions:
        t_type = "📥 Kirim" if tr["type"] == "kirim" else "📤 Chiqim"
        dept_str = tr["department"] if tr["department"] else ("-" if tr["type"] == "kirim" else "Boshqa")
        ws3.append([
            tr["created_at"],
            t_type,
            tr["med_name"],
            f"{tr['quantity']:g} {tr['med_unit']}",
            dept_str,
            tr["comment"] or "-",
            tr["user_name"] or "-",
            tr["id"]
        ])

        type_cell = ws3.cell(row=row_num3, column=2)
        if tr["type"] == "kirim":
            type_cell.font = Font(name="Calibri", size=11, bold=True, color="0070C0")
        else:
            type_cell.font = Font(name="Calibri", size=11, bold=True, color="C00000")

        for c in range(1, 9):
            ws3.cell(row=row_num3, column=c).border = thin_border
            if c in (1, 2, 4, 5, 7, 8):
                ws3.cell(row=row_num3, column=c).alignment = Alignment(horizontal="center")

        row_num3 += 1

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws3.column_dimensions[col_letter].width = max(max_len + 4, 14)

    wb.save(output_filename)
    return output_filename
